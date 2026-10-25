import csv
import os
from copy import copy
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_FILE = os.path.join(BASE_DIR, "data", "arduino_data.csv")
DASHBOARD_FILE = os.path.join(BASE_DIR, "dashboard", "Smart_Security_Auto_Dashboard.xlsx")


HEADERS = [
    "Date",
    "Time",
    "Access",
    "Temperature_C",
    "Humidity_percent",
    "Distance_cm",
]


def read_data():
    if not os.path.exists(CSV_FILE):
        return []

    with open(CSV_FILE, "r", newline="", encoding="utf-8") as file:
        rows = list(csv.reader(file))

    if not rows:
        return []

    data = []
    for row in rows[1:]:
        if len(row) < 6:
            continue
        data.append(row[:6])
    return data


def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def build_dashboard(data):
    os.makedirs(os.path.dirname(DASHBOARD_FILE), exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in data:
        sheet.append(row)

    widths = [14, 12, 14, 18, 20, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    dashboard = workbook.create_sheet("Dashboard")
    dashboard["A1"] = "Smart Security & Environment Monitoring"
    dashboard["A1"].font = Font(bold=True, size=16)

    total = len(data)
    granted = sum(1 for row in data if len(row) >= 3 and row[2] == "GRANTED")
    denied = sum(1 for row in data if len(row) >= 3 and row[2] == "DENIED")

    dashboard["A3"] = "Total Records"
    dashboard["B3"] = total
    dashboard["A4"] = "Granted"
    dashboard["B4"] = granted
    dashboard["A5"] = "Denied"
    dashboard["B5"] = denied

    latest = data[-1] if data else None
    dashboard["A7"] = "Latest Temperature (°C)"
    dashboard["B7"] = number(latest[3]) if latest and len(latest) > 3 else ""
    dashboard["A8"] = "Latest Humidity (%)"
    dashboard["B8"] = number(latest[4]) if latest and len(latest) > 4 else ""
    dashboard["A9"] = "Latest Distance (cm)"
    dashboard["B9"] = number(latest[5]) if latest and len(latest) > 5 else ""

    charts_data = workbook.create_sheet("ChartsData")
    charts_data.append(["Access", "Count"])
    charts_data.append(["GRANTED", granted])
    charts_data.append(["DENIED", denied])

    pie = PieChart()
    pie.title = "Access Results"
    pie.add_data(
        Reference(charts_data, min_col=2, min_row=1, max_row=3),
        titles_from_data=True,
    )
    pie.set_categories(Reference(charts_data, min_col=1, min_row=2, max_row=3))
    dashboard.add_chart(pie, "D3")

    sensor_rows = []
    for row in data:
        if len(row) < 6:
            continue
        temp = number(row[3])
        humidity = number(row[4])
        distance = number(row[5])
        if temp is not None or humidity is not None or distance is not None:
            sensor_rows.append([row[0], row[1], temp, humidity, distance])

    if sensor_rows:
        charts_data["D1"] = "Date"
        charts_data["E1"] = "Time"
        charts_data["F1"] = "Temperature_C"
        charts_data["G1"] = "Humidity_percent"
        charts_data["H1"] = "Distance_cm"
        for row in sensor_rows:
            charts_data.append(row)

        line = LineChart()
        line.title = "Sensor Readings"
        line.y_axis.title = "Value"
        line.x_axis.title = "Reading"
        line.add_data(
            Reference(charts_data, min_col=6, max_col=8, min_row=1, max_row=1 + len(sensor_rows)),
            titles_from_data=True,
        )
        line.set_categories(
            Reference(charts_data, min_col=4, min_row=2, max_row=1 + len(sensor_rows))
        )
        dashboard.add_chart(line, "D18")

    for ws in [dashboard, charts_data]:
        for column in range(1, 10):
            ws.column_dimensions[chr(64 + column)].width = 22

    charts_data.sheet_state = "hidden"
    workbook.save(DASHBOARD_FILE)


def main():
    data = read_data()
    build_dashboard(data)
    print("Dashboard updated successfully!")
    print(f"Open: {DASHBOARD_FILE}")
    print(f"CSV used: {CSV_FILE}")
    print(f"Records: {len(data)}")
    print(f"Granted: {sum(1 for row in data if len(row) >= 3 and row[2] == 'GRANTED')}")
    print(f"Denied: {sum(1 for row in data if len(row) >= 3 and row[2] == 'DENIED')}")


if __name__ == "__main__":
    main()
